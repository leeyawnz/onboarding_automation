import os
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

# input statements
wb_name = input('Enter Excelsheet Name (.xlsx): ')
data_tb_len = input('Enter length of data (Last cell + 1): ')
data_tb_len = int(data_tb_len)
appname = input('Enter Application Team Name: ')

# loading workbook
wb = load_workbook(wb_name)
ws = wb.active

# Query template list
full_template_list = []
for col in range(1, 2):
    char = get_column_letter(col)
    for row in range(1, data_tb_len):
        cell = ws[char + str(row)].value
        full_template_list.append(cell)

i = 0
while i < len(full_template_list):
    with open(f'{appname}.txt', 'a') as file:
        yaml_template = f"""
# {i + 1}.
- pkstatus: true
  dataSource: hasura
  querytemplate: '{full_template_list[i]}'
  pk:
  appid: {appname}
  paramValueSchema:
  required:
"""
        file.write(yaml_template)
        i += 1

txt_file = os.path.join('.', f'{appname}.txt')
yaml_file = txt_file.replace('.txt', '.yaml')
os.rename(txt_file, yaml_file)
